import io
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core import agente, datos, excel, interprete, simulador, tarifario  # noqa: E402


@pytest.fixture(scope="session", autouse=True)
def carpeta_temporal(tmp_path_factory):
    import os

    from core import almacen

    # Las pruebas nunca escriben en el historial ni en la memoria compartidos.
    os.environ["AGENTE_SIN_BACKEND"] = "1"

    almacen.CARPETA = tmp_path_factory.mktemp("data")


@pytest.fixture(scope="session")
def df():
    return datos.cargar_agregado()


@pytest.fixture(scope="session")
def tar():
    return tarifario.cargar()


@pytest.fixture(scope="session")
def cat(df):
    return dict(zip(df["unidad_cod"], df["unidad"]))


# --- Intérprete ----------------------------------------------------------------------------

@pytest.mark.parametrize("texto, esperado", [
    ("¿Cómo impacta en el total de ingresos una rebaja de 5 dólares del derecho de embarque?",
     dict(tipo="simulacion", concepto="Derecho de Embarque", cambio_tipo="absoluto", cambio_valor=-5, cambio_moneda="USD",
          unidades=[])),
    ("como impacta en el total de derecho de embarque de punta arenas una rebaja de 5 dolares aplicado solamente a "
     "punta arenas para fomentar el turismo", dict(unidades=[4101], cambio_valor=-5, region="Magallanes")),
    ("rebaja de US$5 en Punta Arenas con elasticidad 0,3 y dólar a 950",
     dict(elasticidad=0.3, tipo_cambio=950, cambio_moneda="USD")),
    ("Sube 10% las tasas de aterrizaje internacional entre marzo y agosto",
     dict(concepto="Tasas", filtro_descripcion="Aterrizajes", ambito="internacional", meses=[3, 4, 5, 6, 7, 8],
          cambio_tipo="porcentaje", cambio_valor=10)),
    ("¿Cuánto representa Calama en el total de concesiones?", dict(tipo="consulta", concepto="Concesiones",
                                                                  unidades=[1401])),
    ("eliminar el derecho de embarque nacional en la región de Aysén el primer semestre",
     dict(cambio_valor=-100, ambito="nacional", region="Aysén", meses=[1, 2, 3, 4, 5, 6])),
    ("rebaja de 1.000 pesos en el embarque nacional de Isla de Pascua para todo el año",
     dict(cambio_valor=-1000, cambio_moneda="CLP", unidades=[2301], anualizar=True)),
    ("si el embarque internacional baja a 20 dólares en Santiago y los pasajeros aumentan 4%",
     dict(cambio_tipo="nuevo_valor", cambio_valor=20, cambio_moneda="USD", var_pasajeros=4, unidades=[5201])),
    ("aumento de 3% en concesiones de Santiago", dict(concepto="Concesiones", cambio_valor=3, unidades=[5201])),
    ("alza de 2 dólares en el derecho de embarque internacional en el segundo semestre",
     dict(cambio_valor=2, ambito="internacional", meses=[7, 8, 9, 10, 11, 12])),
    ("rebaja de 15% del embarque en Temuco y Valdivia", dict(cambio_tipo="porcentaje", cambio_valor=-15,
                                                             unidades=[3801, 3401])),
    ("ingresos por combustible de antofagasta", dict(tipo="consulta", concepto="Combustible", unidades=[1301])),
])
def test_interprete(texto, esperado, cat):
    c = interprete.interpretar(texto, cat, 2026)
    for k, v in esperado.items():
        real = getattr(c, k)
        if isinstance(v, list):
            assert sorted(real) == sorted(v), (k, real)
        else:
            assert real == v, (k, real)


def test_interprete_dato(cat):
    c = interprete.interpretar("el PIB regional de Magallanes 2024 es 5.900 miles de millones", cat, 2026)
    assert c.tipo == "dato"
    assert c.dato["indicador"] == "pib_regional" and c.dato["region"] == "Magallanes"
    assert c.dato["periodo"] == "2024" and c.dato["valor"] == 5900


# --- Datos y tarifas -------------------------------------------------------------------------

def test_agregado_sin_columnas_sensibles(df):
    assert not {"RUT_", "CLIENTE_", "NUMERO_", "DETALLE_", "rut", "cliente"} & set(df.columns)
    assert set(df["concepto"]) == set(datos.CONCEPTOS)


def test_tarifas_trimestrales(tar):
    assert tarifario.tarifa_mes(tar, "NAC_1C", 2026, 3) == 7814
    assert tarifario.tarifa_mes(tar, "NAC_2C", 2026, 6) == 6039
    # Abril: días 1-9 con la tarifa anterior.
    assert tarifario.tarifa_mes(tar, "NAC_1C", 2026, 4) == pytest.approx((9 * 7814 + 21 * 7923) / 30)
    assert tarifario.tarifa_mes(tar, "INT_MAS500", 2026, 8) == 25
    assert tarifario.tarifa_mes(tar, "INT_MENOS500", 2026, 8) == 8034


def test_interpretar_pagina_dgac():
    html = ("<li>Aeropuertos y Aeródromos de Primera Categoría » $ <strong>8.122</strong></li>"
            "<li>Aeródromos de Segunda Categoría » $ <strong>6.191</strong></li>"
            "<li>Pasajeros que viajen a menos de 270 Km. » $ <strong>3.210</strong></li>"
            "<p>destino a puntos situados a más de 500 kilómetros del Aeropuerto de embarque US$ 25.</p>")
    assert tarifario.interpretar_pagina(html) == {"NAC_1C": 8122, "NAC_2C": 6191, "NAC_270": 3210, "INT_MAS500": 25}


# --- Simulador ---------------------------------------------------------------------------------

def test_base_reproduce_totales(df, tar, cat):
    c = interprete.interpretar("ingresos de derecho de embarque", cat, 2026)
    r = simulador.ejecutar(df, c, tar)
    assert r.niveles.iloc[0]["Base"] == pytest.approx(df.loc[df["concepto"] == "Derecho de Embarque", "monto"].sum())
    assert r.niveles.iloc[-1]["Base"] == pytest.approx(df["monto"].sum())


def test_pasajeros_enteros_en_meses_sin_cambio_tarifario(df, tar):
    """En meses sin reajuste, los pasajeros nacionales inferidos son números enteros (tarifa correcta).

    No es el 100%: algunas filas agregadas suman documentos rezagados emitidos con la tarifa anterior.
    """
    d = simulador._pasajeros(df[(df["mes"].isin([2, 3, 5, 6, 8, 9])) & df["tasa"].isin([21101, 21102, 21103])
                                & (df["monto"] > 0)], tar)
    fraccion = (d["pax"] - d["pax"].round()).abs()
    assert (fraccion < 1e-6).mean() > 0.88


def test_rebaja_usd_punta_arenas_calculo_manual(df, tar, cat):
    tc = 950.0
    c = interprete.interpretar("rebaja de 5 dólares del embarque solo en Punta Arenas", cat, 2026)
    r = simulador.ejecutar(df, c, tar, tc=tc)
    sel = df[(df["unidad_cod"] == 4101) & (df["concepto"] == "Derecho de Embarque")]
    tcm = datos.tipo_cambio_mensual(df)
    esperado = 0.0
    for x in sel.itertuples():
        clase = tarifario.clase_de(x.tasa)
        t = tarifario.tarifa_mes(tar, clase, x.anio, x.mes)
        if tarifario.moneda(tar, clase) == "USD":
            pax = x.monto_usd / t if x.monto_usd > 0 else x.monto / (t * tcm[(x.anio, x.mes)])
            tc_fila = x.monto / x.monto_usd if x.monto_usd > 0 else tcm[(x.anio, x.mes)]
            esperado -= pax * 5 * tc_fila
        else:
            esperado -= x.monto / t * min(5 * tc, t)
    assert r.impacto == pytest.approx(esperado, rel=1e-9)
    # El impacto es el mismo en todos los niveles.
    assert (r.niveles["Diferencia"] - r.impacto).abs().max() < 1


def test_elasticidad_aumenta_pasajeros(df, tar, cat):
    c0 = interprete.interpretar("rebaja de 10% del embarque en Punta Arenas", cat, 2026)
    c1 = interprete.interpretar("rebaja de 10% del embarque en Punta Arenas con elasticidad 0,5", cat, 2026)
    r0, r1 = simulador.ejecutar(df, c0, tar), simulador.ejecutar(df, c1, tar)
    assert r0.impacto == pytest.approx(-0.1 * r0.niveles.iloc[0]["Base"], rel=1e-6)
    assert r1.pasajeros_escenario == pytest.approx(r1.pasajeros_base * 1.05, rel=1e-6)
    assert r1.impacto > r0.impacto


def test_porcentaje_otro_concepto(df, tar, cat):
    c = interprete.interpretar("alza de 10% en concesiones", cat, 2026)
    r = simulador.ejecutar(df, c, tar)
    assert r.impacto == pytest.approx(0.1 * df.loc[df["concepto"] == "Concesiones", "monto"].sum())


def test_anualizar_proyecta_meses(df, tar, cat):
    c = interprete.interpretar("rebaja de 10% del embarque nacional en Calama en el año completo", cat, 2026)
    r = simulador.ejecutar(df, c, tar)
    assert r.meses_proyectados == [10, 11, 12]
    assert len(r.por_mes) == 12


# --- Agente y Excel --------------------------------------------------------------------------

def test_agente_pide_pib_y_lo_recuerda(df, tar):
    r = agente.responder("rebaja de 5 dólares del embarque en Punta Arenas", df, tar, tc_usuario=950)
    assert any(f.indicador == "pib_regional" and f.region == "Magallanes" for f in r.faltantes)
    msg = agente.responder("el PIB regional de Magallanes 2024 es 5.900 miles de millones", df, tar)
    assert "Guardé" in msg.mensaje
    r2 = agente.responder("rebaja de 5 dólares del embarque en Punta Arenas", df, tar, tc_usuario=950)
    assert not r2.faltantes
    assert any("PIB regional" in p for p in r2.resena)


def test_falta_tipo_de_cambio(df, tar, monkeypatch):
    from core import indicadores

    monkeypatch.setattr(indicadores, "obtener", lambda *a, **k: None)
    r = agente.responder("rebaja de 5 dólares del embarque nacional", df, tar)
    assert r.resultado is None and r.faltantes[0].indicador == "dolar" and r.faltantes[0].obligatorio


def test_excel_generado(df, tar):
    r = agente.responder("rebaja de 5 dólares del embarque en Punta Arenas con elasticidad 0,3", df, tar,
                         tc_usuario=950)
    wb = openpyxl.load_workbook(io.BytesIO(excel.generar(r)))
    assert wb.sheetnames == ["Resumen", "Cálculo", "Por unidad", "Por mes", "Gráficos", "Indicadores y fuentes"]
    assert wb["Gráficos"]._charts
    ws = wb["Cálculo"]
    assert ws["B3"].value == 950 and ws["B7"].value == 0.3
    assert str(ws["I12"].value).startswith("=")   # pasajeros con fórmula
    wb2 = openpyxl.load_workbook(io.BytesIO(excel.historial([{"pregunta": "x", "resumen": "y"}])))
    assert wb2["Historial"]["C4"].value == "x"


# --- Tráfico aéreo (JAC / OACI) ------------------------------------------------------------------

def test_interprete_trafico(cat):
    c = interprete.interpretar("¿Cómo va el tráfico aéreo según la JAC?", cat, 2026)
    assert c.tipo == "trafico" and c.fuente == "jac"
    c = interprete.interpretar("Pasajeros embarcados en Calama según la JAC", cat, 2026)
    assert c.tipo == "trafico" and c.unidades == [1401]
    c = interprete.interpretar("datos de la OACI para Latinoamérica", cat, 2026)
    assert c.tipo == "trafico" and c.fuente == "oaci"
    # Una simulación que menciona pasajeros sigue siendo simulación.
    assert interprete.interpretar("rebaja de 5 dólares y los pasajeros aumentan 3%", cat, 2026).tipo == "simulacion"


def test_jac_informe_texto():
    from core import jac

    pagina = ("INFORME MENSUAL DE TRÁFICO AÉREO EN CHILE \nSEPTIEMBRE 2026 \n\n"
              "En septiembre de 2026 se transportaron 2,23 millones de pasajeros en vuelos domésticos e internacionales.\n\n"
              "Mes Dom Int Total Dom Int Total\n"
              "ago-26 1.167.933 1.020.137 2.188.070 -1,5 -2,8 -2,1 \n"
              "sept-26 1.288.851 941.443 2.230.294 1,6 -1,6 0,2 \n"
              "Acum. 2026 11.660.572 9.216.451 20.877.023 -2,8 -0,7 -1,9 \n"
              "Mes Dom Int Total Dom Int Total\n"
              "ago-26 2.759 30.371 33.130 -3,0 3,7 3,1\n")
    rutas = ("Pasajeros en vuelos domésticos por ruta (diez mayores niveles)\n"
             "Ruta Septiembre 2025 Septiembre 2026 Var 26/25 (%) Participación sep-26 (%)\n"
             "SANTIAGO - CALAMA 187.915  198.681  5,7  15,4  \nTOTAL 1 2 3 4\n")
    inf = jac.interpretar_paginas([pagina, rutas])
    assert inf["periodo"] == "Septiembre 2026"
    assert len(inf["serie"]) == 2 and inf["serie"][-1]["domestico"] == 1288851
    assert inf["serie"][-1]["var_total"] == pytest.approx(0.002)
    assert inf["acumulado"]["total"] == 20877023
    assert inf["rutas"][0] == {"tipo": "nacional", "ruta": "SANTIAGO - CALAMA", "anterior": 187915, "actual": 198681,
                               "variacion": pytest.approx(0.057), "participacion": pytest.approx(0.154)}


def test_jac_embarques_y_estacionalidad():
    from core import jac

    e = jac.embarques()
    assert not e.empty and {"anio", "mes", "aeropuerto", "tipo", "pasajeros", "unidad_cod"} <= set(e.columns)
    assert (e["unidad_cod"] == 4101).any()           # Punta Arenas mapeado a su unidad DGAC
    s = jac.estacionalidad([4101])
    assert s and sum(s.values()) == pytest.approx(12)


def test_proyeccion_usa_estacionalidad(df, tar, cat):
    c = interprete.interpretar("alza de 10% del embarque en Punta Arenas en el año completo", cat, 2026)
    r = simulador.ejecutar(df, c, tar, tc=950)
    pm = r.por_mes.set_index("mes")["monto"]
    assert pm[12] > pm[10]                           # diciembre es temporada alta en Punta Arenas


def test_respuesta_trafico_con_unidades(df, tar):
    r = agente.responder("Pasajeros embarcados en Punta Arenas según la JAC", df, tar)
    assert r.tablas and r.grafico and "JAC" in r.resena[0]
    wb = openpyxl.load_workbook(io.BytesIO(excel.generar_tablas(r)))
    assert len(wb.sheetnames) == 1 + len(r.tablas)


def test_oaci_texto():
    from core import oaci

    o = {"periodo": "2026-08", "regiones": {"Latinoamérica/Caribe": {"ask_total": 0.068, "ask_internacional": 0.06}}}
    assert oaci.texto_region(o) == "OACI (08/2026): en Latinoamérica/Caribe, capacidad total (ASK) +6,8%, internacional +6,0% interanual."


# --- Cambio sin magnitud y efecto en pasajeros --------------------------------------------------

PREGUNTA_EL_LOA = ("existe la intencion de bajar la tasa de aterrizaje en el aeródromo el loa, como afecta esta "
                   "situación al total de ingresos por tasas de la Unidad? existe algun efecto en la cantidad de pasajeros?")


def test_cambio_sin_magnitud_da_sensibilidad(df, tar, cat):
    c = interprete.interpretar(PREGUNTA_EL_LOA, cat, 2026)
    assert c.tipo == "simulacion" and c.sensibilidad and c.cambio_valor == -10
    assert c.concepto == "Tasas" and c.filtro_descripcion == "Aterrizajes" and c.unidades == [1401]
    assert c.efecto_pasajeros
    r = agente.responder(PREGUNTA_EL_LOA, df, tar, tc_usuario=950)
    assert r.resultado.impacto == pytest.approx(-0.1 * r.resultado.niveles.iloc[0]["Base"])
    sens = r.tablas["Sensibilidad"]
    assert list(sens["Variación de la tasa"]) == [-0.05, -0.10, -0.15, -0.20, -0.25]
    assert any("no tiene un efecto directo" in p for p in r.resena)
    assert any(f.indicador == "tarifa_pasaje" for f in r.faltantes)


def test_efecto_pasajeros_con_tarifa(df, tar):
    from core import indicadores

    indicadores.guardar_en_memoria(indicadores.Valor(indicador="tarifa_pasaje|Antofagasta", periodo="2026",
                                                     valor=60000, unidad="CLP", fuente="prueba"))
    r = agente.responder(PREGUNTA_EL_LOA, df, tar, tc_usuario=950)
    t = r.tablas["Efecto en pasajeros"]
    dif = r.resultado.impacto
    assert list(t["Elasticidad precio del pasaje"]) == [0.5, 1.0, 1.5]
    assert (t["Pasajeros adicionales"] > 0).all()                    # una rebaja aumenta pasajeros
    assert (t["Efecto neto (tasa + embarque)"] > dif).all()           # el embarque adicional compensa en parte
    wb = openpyxl.load_workbook(io.BytesIO(excel.generar(r)))
    assert {"Sensibilidad", "Efecto en pasajeros"} <= set(wb.sheetnames)


def test_consulta_sin_intencion_sigue_siendo_consulta(cat):
    assert interprete.interpretar("¿Cuánto representa Calama en el total de concesiones?", cat, 2026).tipo == "consulta"
    assert interprete.interpretar("ingresos por tasas de aterrizaje de Calama", cat, 2026).tipo == "consulta"


def test_borrar_historial(df, tar):
    from core import historial

    r = agente.responder("alza de 10% en concesiones", df, tar)
    historial.guardar(r, "prueba")
    assert any(x["id"] == r.id for x in historial.listar()[0])
    assert historial.borrar(r.id)
    assert not any(x["id"] == r.id for x in historial.listar()[0])
